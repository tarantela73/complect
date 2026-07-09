"""Tests for contracts report query and Excel export endpoint (U3)."""

from __future__ import annotations

import uuid
from datetime import date, timezone as tz
from io import BytesIO

import pytest
from openpyxl import load_workbook

from apps.common.tenancy.context import set_tenant_context
from apps.common.tenancy.context_core import RequestContext, set_context
from apps.contragents.models.contragent import Contragent, ContragentType
from apps.contracts.models.contract import Contract, ResourceSupplyType
from apps.reports.queries.contract_report_query import get_contracts_report_rows
from apps.reports.utils.excel import make_excel_response


def _ctx(tenant_id: uuid.UUID) -> None:
    set_context(
        RequestContext(
            tenant_id=tenant_id,
            user_id=uuid.uuid4(),
            request_id=uuid.uuid4(),
            correlation_id=uuid.uuid4(),
        )
    )
    set_tenant_context(tenant_id)


def _make_contragent(tenant_id: uuid.UUID, name: str = "ООО Тест", inn: str = "7700000001") -> Contragent:
    ct = Contragent(
        name=name,
        inn=inn,
        contragent_type=ContragentType.LEGAL_ENTITY,
        tenant_id=tenant_id,
    )
    ct.save()
    return ct


def _make_contract(
    tenant_id: uuid.UUID,
    contragent_id: uuid.UUID | None = None,
    status: str = "ACTIVE",
    total_amount: int = 100_000,
    paid_amount: int = 0,
) -> Contract:
    c = Contract(
        contract_number=f"C-{uuid.uuid4().hex[:6].upper()}",
        title="Test Contract",
        status=status,
        resource_type=ResourceSupplyType.ELECTRICITY,
        total_amount=total_amount,
        paid_amount=paid_amount,
        tenant_id=tenant_id,
        contragent_id=contragent_id,
    )
    c.save()
    return c


@pytest.mark.django_db
class TestContractReportQuery:
    @pytest.fixture(autouse=True)
    def setup(self):
        self.tenant_id = uuid.uuid4()
        _ctx(self.tenant_id)

    def test_returns_all_tenant_contracts(self):
        ct = _make_contragent(self.tenant_id)
        _make_contract(self.tenant_id, contragent_id=ct.id)
        _make_contract(self.tenant_id, contragent_id=ct.id)

        rows = get_contracts_report_rows()
        assert len(rows) == 2

    def test_status_filter(self):
        ct = _make_contragent(self.tenant_id)
        _make_contract(self.tenant_id, contragent_id=ct.id, status="ACTIVE")
        _make_contract(self.tenant_id, contragent_id=ct.id, status="DRAFT")

        rows = get_contracts_report_rows(status="ACTIVE")
        assert len(rows) == 1
        assert rows[0]["status"] == "ACTIVE"

    def test_contragent_join_populates_name_and_inn(self):
        ct = _make_contragent(self.tenant_id, name="ООО Ромашка", inn="7700000099")
        _make_contract(self.tenant_id, contragent_id=ct.id)

        rows = get_contracts_report_rows()
        assert rows[0]["contragent_name"] == "ООО Ромашка"
        assert rows[0]["contragent_inn"] == "7700000099"

    def test_no_contragent_returns_nulls(self):
        _make_contract(self.tenant_id, contragent_id=None)

        rows = get_contracts_report_rows()
        assert rows[0]["contragent_name"] is None
        assert rows[0]["contragent_inn"] is None

    def test_tenant_isolation(self):
        other_tenant = uuid.uuid4()
        ct_own = _make_contragent(self.tenant_id)
        ct_other = _make_contragent(other_tenant, inn="1234567890")
        _make_contract(self.tenant_id, contragent_id=ct_own.id)
        _make_contract(other_tenant, contragent_id=ct_other.id)

        rows = get_contracts_report_rows()
        assert len(rows) == 1

    def test_empty_result_returns_empty_list(self):
        rows = get_contracts_report_rows()
        assert rows == []

    def test_amounts_are_raw_kopecks(self):
        ct = _make_contragent(self.tenant_id)
        _make_contract(self.tenant_id, contragent_id=ct.id, total_amount=500_00, paid_amount=200_00)

        rows = get_contracts_report_rows()
        assert rows[0]["total_amount"] == 500_00
        assert rows[0]["paid_amount"] == 200_00
        assert rows[0]["remaining_amount"] == 300_00


@pytest.mark.django_db
class TestContractExcelOutput:
    @pytest.fixture(autouse=True)
    def setup(self):
        self.tenant_id = uuid.uuid4()
        _ctx(self.tenant_id)

    def test_excel_response_headers(self):
        from openpyxl import Workbook

        wb = Workbook()
        resp = make_excel_response(wb, "contracts_report.xlsx")
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in resp["Content-Type"]
        assert 'filename="contracts_report.xlsx"' in resp["Content-Disposition"]

    def test_workbook_has_header_row(self):
        ct = _make_contragent(self.tenant_id)
        _make_contract(self.tenant_id, contragent_id=ct.id)

        rows = get_contracts_report_rows()

        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        headers = [
            "№",
            "Номер договора",
            "Наименование",
            "Статус",
            "Вид ресурса",
            "Контрагент",
            "ИНН",
            "Сумма (руб.)",
            "Оплачено (руб.)",
            "Остаток (руб.)",
            "Дата создания",
        ]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)

        resp = make_excel_response(wb, "test.xlsx")
        wb2 = load_workbook(BytesIO(resp.content))
        ws2 = wb2.active
        assert [cell.value for cell in ws2[1]] == headers
