"""Unit tests for reports utility helpers (U2)."""
from __future__ import annotations

import pytest
from openpyxl import Workbook

from apps.reports.utils.excel import make_excel_response
from apps.reports.utils.pdf import make_pdf_response


@pytest.mark.django_db
class TestMakeExcelResponse:
    def test_content_type(self):
        wb = Workbook()
        response = make_excel_response(wb, "report.xlsx")
        assert response["Content-Type"] == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    def test_content_disposition(self):
        wb = Workbook()
        response = make_excel_response(wb, "report.xlsx")
        assert response["Content-Disposition"] == 'attachment; filename="report.xlsx"'

    def test_response_body_is_valid_xlsx(self):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Test"  # type: ignore[index]
        response = make_excel_response(wb, "test.xlsx")
        # Valid XLSX starts with PK zip header
        assert response.content[:2] == b"PK"


class TestMakePdfResponse:
    def test_content_type(self):
        response = make_pdf_response(b"%PDF-fake", "doc.pdf")
        assert response["Content-Type"] == "application/pdf"

    def test_content_disposition(self):
        response = make_pdf_response(b"%PDF-fake", "doc.pdf")
        assert response["Content-Disposition"] == 'attachment; filename="doc.pdf"'

    def test_body_passed_through(self):
        data = b"%PDF-1.4 fake content"
        response = make_pdf_response(data, "x.pdf")
        assert response.content == data
