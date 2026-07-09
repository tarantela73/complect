import uuid
from datetime import date
from typing import Optional

from django.core.exceptions import PermissionDenied
from django.http import HttpRequest
from ninja import Query, Router
from openpyxl import Workbook
from openpyxl.styles import Font

from apps.common.api.errors import make_error_response
from apps.common.api.permissions import require_permission
from apps.reports.pdf_service import render_contract_pdf, render_contragent_pdf
from apps.reports.queries.contract_report_query import get_contracts_report_rows
from apps.reports.queries.finance_report_query import get_finance_report_rows
from apps.reports.queries.payments_report_query import get_payments_report_rows
from apps.reports.utils.excel import make_excel_response
from apps.reports.utils.pdf import make_pdf_response

router = Router(tags=["reports"])


def _amount_rub(kopecks: int) -> float:
    return round(kopecks / 100, 2)


STATUS_LABELS = {
    "DRAFT": "Черновик",
    "ACTIVE": "Активен",
    "LEGAL_HOLD": "Юр. удержание",
    "CLOSED": "Закрыт",
}

RESOURCE_LABELS = {
    "ELECTRICITY": "Электроснабжение",
    "GAS": "Газоснабжение",
    "WATER": "Водоснабжение",
    "HEAT": "Теплоснабжение",
    "OTHER": "Прочее",
}

TRANSACTION_TYPE_LABELS = {
    "PAYMENT": "Оплата",
    "REFUND": "Возврат",
    "CORRECTION": "Корректировка",
}

TRANSACTION_STATUS_LABELS = {
    "PENDING": "Ожидает",
    "POSTED": "Проведена",
    "REVERSED": "Сторно",
    "FAILED": "Ошибка",
}


@router.get("/contracts/export/", response={403: dict}, auth=None)
@require_permission("read", "contracts")
def export_contracts(
    request: HttpRequest,
    status: Optional[str] = Query(None),
    resource_type: Optional[str] = Query(None),
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    contragent_id: Optional[uuid.UUID] = Query(None),
):
    try:
        rows = get_contracts_report_rows(
            status=status,
            resource_type=resource_type,
            date_from=date_from,
            date_to=date_to,
            contragent_id=contragent_id,
        )
    except PermissionDenied:
        return 403, make_error_response("FORBIDDEN", "Нет доступа.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Договоры"

    headers = [
        "№",
        "Номер договора",
        "Наименование",
        "Статус",
        "Вид ресурса",
        "Контрагент",
        "ИНН",
        "Сумма (руб.)",
        "Оплачено (руб.)",
        "Остаток (руб.)",
        "Дата создания",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for i, row in enumerate(rows, start=1):
        ws.append(
            [
                i,
                row["contract_number"],
                row["title"],
                STATUS_LABELS.get(row["status"], row["status"]),
                RESOURCE_LABELS.get(row["resource_type"], row["resource_type"]),
                row["contragent_name"] or "",
                row["contragent_inn"] or "",
                _amount_rub(row["total_amount"]),
                _amount_rub(row["paid_amount"]),
                _amount_rub(row["remaining_amount"]),
                row["created_at"].strftime("%d.%m.%Y") if row["created_at"] else "",
            ]
        )

    return make_excel_response(wb, "contracts_report.xlsx")


@router.get("/finance/export/", response={403: dict}, auth=None)
@require_permission("read", "finance")
def export_finance(
    request: HttpRequest,
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    contract_id: Optional[uuid.UUID] = Query(None),
    transaction_type: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
):
    try:
        rows = get_finance_report_rows(
            date_from=date_from,
            date_to=date_to,
            contract_id=contract_id,
            transaction_type=transaction_type,
            status=status,
        )
    except PermissionDenied:
        return 403, make_error_response("FORBIDDEN", "Нет доступа.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Транзакции"

    headers = [
        "Дата",
        "Номер договора",
        "Тип операции",
        "Статус",
        "Сумма (руб.)",
        "ИНН контрагента",
        "Внешний реф.",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append(
            [
                row["value_date"].strftime("%d.%m.%Y") if row["value_date"] else "",
                row["contract_number"] or "",
                TRANSACTION_TYPE_LABELS.get(row["transaction_type"], row["transaction_type"]),
                TRANSACTION_STATUS_LABELS.get(row["status"], row["status"]),
                _amount_rub(row["amount"]),
                row["counterparty_inn"] or "",
                row["external_ref"] or "",
            ]
        )

    return make_excel_response(wb, "finance_report.xlsx")


MATCH_STATUS_LABELS = {
    "UNMATCHED": "Не сопоставлена",
    "PARTIAL": "Частично",
    "MATCHED": "Сопоставлена",
    "AMBIGUOUS": "Неоднозначно",
}


@router.get("/payments/export/", response={403: dict}, auth=None)
@require_permission("read", "payments")
def export_payments(
    request: HttpRequest,
    statement_id: Optional[uuid.UUID] = Query(None),
    match_status: Optional[str] = Query(None),
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
):
    try:
        rows = get_payments_report_rows(
            statement_id=statement_id,
            match_status=match_status,
            date_from=date_from,
            date_to=date_to,
        )
    except PermissionDenied:
        return 403, make_error_response("FORBIDDEN", "Нет доступа.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Выписка"

    headers = [
        "Дата",
        "Сумма",
        "Направление",
        "Контрагент",
        "ИНН",
        "Описание",
        "Реквизит",
        "Статус матчинга",
        "Уверенность (%)",
        "ID договора",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    total = len(rows)
    matched = sum(1 for r in rows if r["match_status"] == "MATCHED")
    unmatched = total - matched

    for row in rows:
        ws.append(
            [
                row["transaction_date"].strftime("%d.%m.%Y") if row["transaction_date"] else "",
                float(row["amount"]) if row["amount"] is not None else "",
                row["direction"] or "",
                row["counterparty_name"] or "",
                row["counterparty_inn"] or "",
                row["description"] or "",
                row["reference"] or "",
                MATCH_STATUS_LABELS.get(row["match_status"], row["match_status"]),
                row["match_confidence"],
                str(row["contract_id"]) if row["contract_id"] else "",
            ]
        )

    ws.append([])
    ws.append(["Итого:", total, "Сопоставлено:", matched, "Не сопоставлено:", unmatched])

    return make_excel_response(wb, "payments_report.xlsx")


@router.get("/contracts/{contract_id}/pdf/", response={404: dict, 403: dict}, auth=None)
@require_permission("read", "contracts")
def export_contract_pdf(request: HttpRequest, contract_id: uuid.UUID):
    try:
        pdf_bytes = render_contract_pdf(contract_id)
    except PermissionDenied:
        return 404, make_error_response("CONTRACT_NOT_FOUND", "Договор не найден.")
    return make_pdf_response(pdf_bytes, f"contract_{contract_id}.pdf")


@router.get("/contragents/{contragent_id}/pdf/", response={404: dict, 403: dict}, auth=None)
@require_permission("read", "contragents")
def export_contragent_pdf(request: HttpRequest, contragent_id: uuid.UUID):
    try:
        pdf_bytes = render_contragent_pdf(contragent_id)
    except PermissionDenied:
        return 404, make_error_response("CONTRAGENT_NOT_FOUND", "Контрагент не найден.")
    return make_pdf_response(pdf_bytes, f"contragent_{contragent_id}.pdf")
